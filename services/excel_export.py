"""Generate a new, editable XLSX from selected fields and rows without executing formulas."""
from __future__ import annotations

import io
import math
from datetime import date, datetime

MAX_SHEET_DATA_ROWS = 1_048_575
MAX_SHEET_COLUMNS = 16_384


class ExcelContentError(RuntimeError):
    """A cell cannot be represented without silently changing its content."""


def _cell_value(value):
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.isoformat()
    if value is None or isinstance(value, (str, bool, int, float, date, datetime)):
        return value
    return str(value)


def build_workbook(fields, rows, highlighted, format_value):
    """Write in streaming mode; the header is always the first worksheet row."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    # Validate before opening the streaming writer: never truncate a large cell
    # or silently turn source text beginning with '=' into an executable formula.
    for field in fields:
        if len(field) > 32_767 or ILLEGAL_CHARACTERS_RE.search(field):
            raise ExcelContentError(f'列名“{field[:60]}”超出 Excel 文本限制或包含无效控制字符。')
    for row_number, row in enumerate(rows, 2):
        for field in fields:
            value = _cell_value(row.get(field))
            if isinstance(value, str) and (len(value) > 32_767 or ILLEGAL_CHARACTERS_RE.search(value)):
                raise ExcelContentError(f'第 {row_number} 行“{field}”超出 Excel 的 32767 字符限制或包含无效控制字符，已取消发送。')

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('数据')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(fields))}{len(rows) + 1}'
    header_fill = PatternFill('solid', fgColor='385D8A')
    header_font = Font(bold=True, color='FFFFFF')
    marked_fill = PatternFill('solid', fgColor='FFF0CC')
    for index, field in enumerate(fields, 1):
        sample = [len(str(row.get(field) or '')[:100]) for row in rows[:80]]
        sheet.column_dimensions[get_column_letter(index)].width = min(50, max(12, len(field) * 2 + 2, max(sample, default=0) + 2))
    header = []
    for field in fields:
        cell = WriteOnlyCell(sheet, value=field)
        cell.data_type = 's'
        cell.fill, cell.font = header_fill, header_font
        header.append(cell)
    sheet.append(header)
    for row, marked in zip(rows, highlighted):
        cells = []
        for field in fields:
            value, number_format = format_value(field, _cell_value(row.get(field)))
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = 's'
            if number_format:
                cell.number_format = number_format
            if marked:
                cell.fill = marked_fill
            cells.append(cell)
        sheet.append(cells)
    target = io.BytesIO()
    try:
        workbook.save(target)
        return target.getvalue()
    finally:
        workbook.close()
