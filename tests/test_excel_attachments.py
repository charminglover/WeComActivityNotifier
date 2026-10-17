"""Verify workbook contents, size splitting, and file upload/send protocol without live pushes."""
import io
import random
import string
import unittest
from datetime import date
from unittest.mock import patch

from openpyxl import load_workbook
import engine
from engine import EXCEL_MODE, Job, WeComSender, prepare_messages
from services.excel_export import ExcelContentError
from services.notification_workflow import prepare_selection


class TestExcelAttachments(unittest.TestCase):
    def test_editable_values_selected_columns_styles_and_literal_formula_text(self):
        job = Job('CON', ['文本', '金额', '完成率', '日期'], [
            {'文本': '=HYPERLINK("https://example.test")', '金额': 123.5,
             '完成率': 0.42, '日期': date(2026, 10, 17), '未选择': 'exclude'}])
        prepared = prepare_messages([job], EXCEL_MODE, highlight_conditions=[
            {'field': '金额', 'operator': '>', 'value': '100'}])[0]
        self.assertEqual(prepared.filename, '_CON.xlsx')
        workbook = load_workbook(io.BytesIO(prepared.payload))
        sheet = workbook.active
        self.assertEqual([cell.value for cell in sheet[1]], job.fields)
        self.assertEqual(sheet['A2'].data_type, 's')
        self.assertEqual(sheet['A2'].value, job.rows[0]['文本'])
        self.assertEqual(sheet['B2'].value, 123.5)
        self.assertEqual(sheet['C2'].value, 0.42)
        self.assertEqual(sheet['C2'].number_format, '0.00%')
        self.assertEqual(sheet['D2'].value.date(), date(2026, 10, 17))
        self.assertEqual(sheet['A2'].fill.fgColor.rgb[-6:], 'FFF0CC')
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:D2')
        self.assertEqual(sheet.max_column, 4)
        workbook.close()

    def test_combined_mode_attaches_full_table_per_destination(self):
        fields = ['区域'] + [f'Column{i}' for i in range(12)]
        rows = [dict({field: f'{field}-{i}' for field in fields},
                     区域='North' if i < 30 else 'South', _id=str(i)) for i in range(61)]
        messages = prepare_selection({'ids': {str(i) for i in range(60)}, 'records': rows,
            'fields': fields, 'title': '筛选结果', 'group_field': '', 'route_field': '区域',
            'mode': 'Markdown表格', 'include_excel': True})
        attachments = [message for message in messages if message.mode == EXCEL_MODE]
        self.assertEqual(len(attachments), 2)
        self.assertTrue(any(message.mode == 'Markdown表格' for message in messages))
        for attachment in attachments:
            self.assertEqual(len(attachment.job.rows), 30)
            workbook = load_workbook(io.BytesIO(attachment.payload), read_only=True)
            values = list(workbook.active.values)
            self.assertEqual(list(values[0]), fields)
            self.assertEqual(len(values), 31)
            self.assertTrue(all(row[0] == attachment.job.route_value for row in values[1:]))
            self.assertNotIn('Column0-60', {row[1] for row in values[1:]})
            workbook.close()

    def test_over_limit_files_split_and_preserve_routing(self):
        randomizer = random.Random(17)
        rows = [{'区域': 'North', '内容': ''.join(randomizer.choices(string.ascii_letters, k=200)),
                 '_id': str(i)} for i in range(40)]
        jobs = engine.make_jobs(rows, ['区域', '内容'], page_size=None,
                                columns_per_page=None, route_field='区域')
        with patch.object(engine, 'MAX_EXCEL_BYTES', 6500):
            messages = prepare_messages(jobs, EXCEL_MODE)
        self.assertGreater(len(messages), 1)
        self.assertEqual(sum(len(m.job.rows) for m in messages), 40)
        self.assertEqual([m.job.part for m in messages], list(range(1, len(messages) + 1)))
        for message in messages:
            self.assertLessEqual(len(message.payload), 6500)
            self.assertEqual(message.job.route_value, 'North')
            self.assertIn(f'_{message.job.part}-{len(messages)}.xlsx', message.filename)
            workbook = load_workbook(io.BytesIO(message.payload), read_only=True)
            self.assertEqual(len(list(workbook.active.values)), len(message.job.rows) + 1)
            workbook.close()

    def test_unrepresentable_cells_fail_without_truncation(self):
        for text in ('x' * 32768, 'bad\x01text'):
            with self.subTest(text_length=len(text)):
                with self.assertRaises(ExcelContentError):
                    prepare_messages([Job('标题', ['数据'], [{'数据': text}])], EXCEL_MODE)

    def test_robot_upload_uses_recipient_key_and_then_sends_file_media_id(self):
        message = prepare_messages([Job('华东报表', ['数据'], [{'数据': 'test'}])], EXCEL_MODE)[0]
        sender = WeComSender('群机器人Webhook',
            webhook='https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=default')
        target = 'https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=target%2Bkey'
        with patch.object(sender, '_request', side_effect=[{'media_id': 'robot-file'}, {'errcode': 0}]) as request:
            sender.send(message, EXCEL_MODE, recipient=target)
            upload, send = request.call_args_list
            self.assertEqual(upload.args[1], 'https://qyapi.weixin.qq.com/cgi-bin/webhook/upload_media')
            self.assertEqual(upload.kwargs['params'], {'key': 'target+key', 'type': 'file'})
            self.assertIn(f'filelength={len(message.payload)}'.encode(), upload.kwargs['data'])
            self.assertIn(message.filename.encode('utf-8'), upload.kwargs['data'])
            self.assertIn(message.payload, upload.kwargs['data'])
            self.assertEqual(send.args[1], target)
            self.assertEqual(send.kwargs['json'], {'msgtype': 'file', 'file': {'media_id': 'robot-file'}})
        with patch.object(sender, '_request', return_value={}) as request:
            with self.assertRaisesRegex(RuntimeError, 'media_id'):
                sender.send(message, EXCEL_MODE, recipient=target)
            self.assertEqual(request.call_count, 1)

    def test_application_file_upload_and_message(self):
        message = prepare_messages([Job('报表', ['数据'], [{'数据': 'test'}])], EXCEL_MODE)[0]
        sender = WeComSender('自建应用消息', corpid='corp', secret='secret', agentid='100', touser='default')
        with patch.object(sender, '_request', side_effect=[{'access_token': 'token'},
                {'media_id': 'app-file'}, {'errcode': 0}]) as request:
            sender.send(message, EXCEL_MODE, recipient='user-a|user-b')
            upload, send = request.call_args_list[1:]
            self.assertEqual(upload.kwargs['params'], {'access_token': 'token', 'type': 'file'})
            self.assertEqual(send.kwargs['json'], {'touser': 'user-a|user-b', 'agentid': 100,
                'msgtype': 'file', 'file': {'media_id': 'app-file'}})


if __name__ == '__main__':
    unittest.main()
